from fastapi import APIRouter, File, UploadFile, HTTPException
from openpyxl import load_workbook

from app.database import get_connection
from app.auth import hash_password

router = APIRouter()

COLUMNAS_REQUERIDAS = {"Correo", "Clave", "Rol"}
ROL_MAPEO = {"directora": "directora", "docente": "docente"}


@router.post("/seed/")
def seed_usuarios(file: UploadFile = File(...)):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "El archivo debe ser .xlsx")

    wb = load_workbook(file.file, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        raise HTTPException(400, "El archivo Excel está vacío")

    headers = [str(h).strip() if h else "" for h in rows[0]]
    header_set = set(headers)

    faltantes = COLUMNAS_REQUERIDAS - header_set
    if faltantes:
        raise HTTPException(
            400,
            f"Faltan columnas requeridas: {', '.join(sorted(faltantes))}",
        )

    col_idx = {h: i for i, h in enumerate(headers)}
    conn = get_connection()
    cursor = conn.cursor()

    insertados = 0
    for row in rows[1:]:
        correo = str(row[col_idx["Correo"]]).strip() if row[col_idx["Correo"]] else ""
        clave = str(row[col_idx["Clave"]]).strip() if row[col_idx["Clave"]] else ""
        rol_raw = str(row[col_idx["Rol"]]).strip().lower() if row[col_idx["Rol"]] else ""
        nombre = correo.split("@")[0]

        if not correo or not clave:
            continue

        rol = ROL_MAPEO.get(rol_raw)
        if not rol:
            continue

        clave_hash = hash_password(clave)
        try:
            cursor.execute(
                "INSERT INTO usuarios (correo, clave_hash, nombre, rol) VALUES (%s, %s, %s, %s)",
                (correo, clave_hash, nombre, rol),
            )
            insertados += 1
        except Exception:
            conn.rollback()

    conn.commit()
    cursor.close()
    conn.close()
    wb.close()

    return {"mensaje": f"Seed completado", "usuarios_insertados": insertados}
